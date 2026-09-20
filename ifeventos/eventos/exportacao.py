"""Exportação dos relatórios em CSV, XLSX e PDF.

Recebe cabeçalhos + linhas já prontos (a seleção de colunas e os filtros são
resolvidos na view) e devolve uma `HttpResponse` para download.
"""

import csv
import io

from django.http import HttpResponse


def _csv(cabecalhos, linhas):
    buffer = io.StringIO()
    buffer.write("\ufeff")  # BOM: Excel abre os acentos corretamente
    escritor = csv.writer(buffer)
    escritor.writerow(cabecalhos)
    escritor.writerows(linhas)
    return buffer.getvalue().encode("utf-8")


def _xlsx(cabecalhos, linhas, abas=None):
    """Gera o XLSX com uma ou mais abas.

    `abas` é uma lista de `(nome_aba, cabecalhos, linhas)`. Quando informado,
    cada aba é escrita separadamente; sem `abas`, usa uma única aba com
    `cabecalhos`/`linhas` (comportamento anterior).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    planilha = Workbook()
    primeira = True
    for nome_aba, ab_cabecalhos, ab_linhas in (abas or [("Resumo", cabecalhos, linhas)]):
        aba = planilha.active if primeira else planilha.create_sheet()
        primeira = False
        aba.title = nome_aba[:31] or "Planilha"
        aba.append(ab_cabecalhos)
        for celula in aba[1]:
            celula.font = Font(bold=True)
        for linha in ab_linhas:
            aba.append(linha)

        for indice, cabecalho in enumerate(ab_cabecalhos, start=1):
            valores = [len(str(cabecalho))]
            valores += [len(str(linha[indice - 1])) for linha in ab_linhas if len(linha) >= indice]
            largura = min(max(max(valores) + 2, 10), 40)
            aba.column_dimensions[aba.cell(row=1, column=indice).column_letter].width = largura

    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()


def _pdf(cabecalhos, linhas, titulo=None):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    documento = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        leftMargin=10 * mm, rightMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm,
    )
    estilos = getSampleStyleSheet()
    elementos = []
    if titulo:
        elementos.append(Paragraph(titulo, estilos["Title"]))
        elementos.append(Spacer(1, 6))

    dados = [cabecalhos] + [[str(valor) for valor in linha] for linha in linhas]
    tabela = Table(dados, repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2f9e41")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f6f7f8")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    elementos.append(tabela)
    documento.build(elementos)
    return buffer.getvalue()


def exportar(formato, nome_base, cabecalhos, linhas, titulo=None, abas=None):
    """Devolve a `HttpResponse` do formato pedido (csv, xlsx ou pdf).

    `abas` (opcional) é uma lista de `(nome_aba, cabecalhos, linhas)`: no XLSX,
    cada item vira uma aba. CSV/PDF ignoram `abas` (usam só `cabecalhos`/`linhas`).
    """
    if formato == "csv":
        conteudo, content_type, extensao = _csv(cabecalhos, linhas), "text/csv; charset=utf-8", "csv"
    elif formato == "xlsx":
        conteudo = _xlsx(cabecalhos, linhas, abas=abas)
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        extensao = "xlsx"
    elif formato == "pdf":
        conteudo, content_type, extensao = _pdf(cabecalhos, linhas, titulo), "application/pdf", "pdf"
    else:
        raise ValueError("formato de exportação desconhecido: %r" % (formato,))

    resposta = HttpResponse(conteudo, content_type=content_type)
    resposta["Content-Disposition"] = 'attachment; filename="%s.%s"' % (nome_base, extensao)
    return resposta
