"""Testes dos certificados: elegibilidade, render e telas de configuração."""

import io
import os
import shutil
import tempfile
from datetime import date, datetime
from datetime import timezone as tz

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse

from eventos import certificados
from eventos.models import (
    Atividade,
    Certificado,
    ConfiguracaoCertificado,
    Evento,
    Inscricao,
    Presenca,
)

U = get_user_model()
SENHA = "SenhaForte123!"


def _evento(organizador, **kwargs):
    dados = dict(
        title="Evento Cert", description="d", local="Auditório",
        data_inicio=date(2026, 10, 10), data_fim=date(2026, 10, 12),
    )
    dados.update(kwargs)
    return Evento.objects.create(organizador=organizador, **dados)


def _atividade(evento, **kwargs):
    dados = dict(
        titulo="Atividade", descricao="d", local="Sala",
        data_hora_inicio=datetime(2026, 10, 10, 10, 0, tzinfo=tz.utc),
        data_hora_fim=datetime(2026, 10, 10, 11, 0, tzinfo=tz.utc),
        n_vagas=10, emite_certificado=True,
    )
    dados.update(kwargs)
    return Atividade.objects.create(evento=evento, **dados)


class ElegibilidadeTests(TestCase):
    def setUp(self):
        self.dono = U.objects.create_user(email="dono@cert.test", password=SENHA, is_organizador=True)
        self.evento = _evento(self.dono)
        self.pessoa = U.objects.create_user(email="p@cert.test", password=SENHA)

    def test_atividade_sem_emite_certificado_nao_elegivel(self):
        a = _atividade(self.evento, emite_certificado=False)
        Presenca.objects.create(atividade=a, participante=self.pessoa)
        self.assertFalse(certificados.participantes_da_atividade(a).exists())

    def test_presenca_torna_elegivel(self):
        a = _atividade(self.evento)
        Presenca.objects.create(atividade=a, participante=self.pessoa)
        self.assertIn(self.pessoa, certificados.participantes_da_atividade(a))

    def test_inscricao_confirmada_torna_elegivel(self):
        a = _atividade(self.evento)
        Inscricao.objects.create(atividade=a, participante=self.pessoa, confirmada=True)
        self.assertIn(self.pessoa, certificados.participantes_da_atividade(a))

    def test_evento_por_percentual(self):
        a1 = _atividade(self.evento, titulo="A1")
        a2 = _atividade(self.evento, titulo="A2")
        a3 = _atividade(self.evento, titulo="A3")
        a4 = _atividade(self.evento, titulo="A4")
        # 3 de 4 = 75% -> elegível; outro com 2 de 4 = 50% -> não.
        for a in (a1, a2, a3):
            Presenca.objects.create(atividade=a, participante=self.pessoa)
        outro = U.objects.create_user(email="q@cert.test", password=SENHA)
        for a in (a1, a2):
            Presenca.objects.create(atividade=a, participante=outro)
        elegiveis = set(certificados.participantes_do_evento(self.evento))
        self.assertIn(self.pessoa, elegiveis)
        self.assertNotIn(outro, elegiveis)

    def test_limiar_configuravel(self):
        self.evento.percentual_certificado = 50
        self.evento.save(update_fields=["percentual_certificado"])
        a1 = _atividade(self.evento, titulo="A1")
        _atividade(self.evento, titulo="A2")
        Presenca.objects.create(atividade=a1, participante=self.pessoa)
        self.assertIn(self.pessoa, certificados.participantes_do_evento(self.evento))


class RenderTests(TestCase):
    def setUp(self):
        self.dono = U.objects.create_user(email="dono2@cert.test", password=SENHA, is_organizador=True)
        self.evento = _evento(self.dono)

    def test_render_pdf_gera_pdf_valido(self):
        config = ConfiguracaoCertificado.objects.create(
            evento=self.evento, titulo="CERTIFICADO",
            corpo="Certificamos que {{nome}} participou de {{atividade}} ({{carga_horaria}}).",
        )
        contexto = certificados.contexto_certificado(self.dono, evento=self.evento, config=config)
        dados = certificados.render_pdf(contexto, config)
        self.assertTrue(dados.startswith(b"%PDF"))

    def test_substituir_variaveis(self):
        texto = certificados.substituir_variaveis("Oi {{nome}}!", {"nome": "Ana"})
        self.assertEqual(texto, "Oi Ana!")


class ConfigViewTests(TestCase):
    def setUp(self):
        self.dono = U.objects.create_user(email="dono3@cert.test", password=SENHA, is_organizador=True)
        self.evento = _evento(self.dono)
        self.participante = U.objects.create_user(email="x@cert.test", password=SENHA)

    def test_participante_recebe_403(self):
        self.client.force_login(self.participante)
        r = self.client.get(
            reverse("organizador:certificado_config", args=[self.evento.id, "evento"])
        )
        self.assertEqual(r.status_code, 403)

    def test_dono_abre_e_salva(self):
        from eventos.models import Assinante

        assinante = Assinante.objects.create(nome="Diretor Geral", cargo="Diretor")
        self.client.force_login(self.dono)
        url = reverse("organizador:certificado_config", args=[self.evento.id, "evento"])
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200)
        r = self.client.post(url, {
            "titulo": "CERTIFICADO DE PARTICIPAÇÃO",
            "corpo": "Certificamos que {{nome}}.",
            "modo_layout": "texto",
            "enviar_email": "on",
            "assinantes_escolhidos": [assinante.id],
        })
        self.assertEqual(r.status_code, 302)
        config = ConfiguracaoCertificado.objects.get(
            evento=self.evento, escopo="evento", atividade__isnull=True
        )
        self.assertEqual(config.assinaturas.count(), 1)
        self.assertEqual(config.assinaturas.first().nome, "Diretor Geral")

    def test_preview_retorna_pdf(self):
        self.client.force_login(self.dono)
        r = self.client.get(reverse("organizador:certificado_preview", args=[self.evento.id]))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/pdf")
        self.assertTrue(r.content.startswith(b"%PDF"))


class EmissaoTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls._media = tempfile.mkdtemp(prefix="media_cert_")
        cls.enterClassContext(override_settings(MEDIA_ROOT=cls._media))
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(cls._media, ignore_errors=True)

    def setUp(self):
        from django.core import mail

        mail.outbox = []
        self.dono = U.objects.create_user(
            email="dono4@cert.test", password=SENHA, is_organizador=True
        )
        self.evento = _evento(self.dono)
        self.pessoa = U.objects.create_user(email="aluno@cert.test", password=SENHA)
        self.atividade = _atividade(self.evento)
        Presenca.objects.create(atividade=self.atividade, participante=self.pessoa)

    def test_emitir_atividade_idempotente(self):
        certificado, criado = certificados.emitir(self.pessoa, atividade=self.atividade)
        self.assertTrue(criado)
        self.assertEqual(certificado.tipo, "atividade")
        certificado2, criado2 = certificados.emitir(self.pessoa, atividade=self.atividade)
        self.assertFalse(criado2)
        self.assertEqual(certificado2.pk, certificado.pk)
        self.assertEqual(Certificado.objects.count(), 1)

    def test_emitir_evento_tipo_evento(self):
        certificado, criado = certificados.emitir(self.pessoa, evento=self.evento)
        self.assertTrue(criado)
        self.assertEqual(certificado.tipo, "evento")

    def test_email_enviado_quando_configurado(self):
        from django.core import mail

        ConfiguracaoCertificado.objects.create(
            evento=self.evento, escopo="atividade", enviar_email=True
        )
        certificados.emitir(self.pessoa, atividade=self.atividade)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].attachments[0][0], "certificado.pdf")

    def test_email_nao_enviado_quando_desligado(self):
        from django.core import mail

        ConfiguracaoCertificado.objects.create(
            evento=self.evento, escopo="atividade", enviar_email=False
        )
        certificados.emitir(self.pessoa, atividade=self.atividade)
        self.assertEqual(len(mail.outbox), 0)

    def test_view_emitir_certificados_atividade(self):
        self.client.force_login(self.dono)
        r = self.client.post(
            reverse("organizador:emitir_certificados_atividade", args=[self.atividade.id])
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(Certificado.objects.count(), 1)

    def test_emissao_docx_sem_libreoffice_relata_falha(self):
        from unittest import mock

        from django.core.files.base import ContentFile
        from docx import Document

        buffer = io.BytesIO()
        Document().save(buffer)
        config = ConfiguracaoCertificado.objects.create(
            evento=self.evento, escopo="atividade", modo_layout="docx"
        )
        config.template_docx.save("m.docx", ContentFile(buffer.getvalue()), save=True)

        self.client.force_login(self.dono)
        with mock.patch.object(
            certificados, "_docx_para_pdf", side_effect=FileNotFoundError("libreoffice")
        ):
            r = self.client.post(
                reverse("organizador:emitir_certificados_atividade", args=[self.atividade.id])
            )
        self.assertEqual(r.status_code, 200)
        dados = r.json()
        self.assertEqual(dados["certificados"], [])
        self.assertTrue(dados.get("falhas"))
        self.assertEqual(Certificado.objects.count(), 0)

    def test_view_emitir_certificado_do_evento(self):
        self.client.force_login(self.dono)
        r = self.client.post(
            reverse("organizador:emitir_certificados_evento", args=[self.evento.id])
        )
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["certificados"], [self.pessoa.id])
        self.assertTrue(
            Certificado.objects.filter(
                participante=self.pessoa, evento=self.evento, tipo="evento"
            ).exists()
        )

    def test_view_emitir_certificados_de_todas_as_atividades(self):
        self.client.force_login(self.dono)
        r = self.client.post(
            reverse("organizador:emitir_certificados_todas_atividades", args=[self.evento.id])
        )
        self.assertEqual(r.status_code, 200)
        self.assertIn(self.pessoa.id, r.json()["certificados"])
        self.assertTrue(
            Certificado.objects.filter(
                participante=self.pessoa, atividade=self.atividade
            ).exists()
        )


def _docx_template_bytes(texto="Certificado de {{ nome }}"):
    from docx import Document

    documento = Document()
    documento.add_paragraph(texto)
    buffer = io.BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


class DocxRenderTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls._media = tempfile.mkdtemp(prefix="media_docx_")
        cls.enterClassContext(override_settings(MEDIA_ROOT=cls._media))
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(cls._media, ignore_errors=True)

    def setUp(self):
        from django.core.files.base import ContentFile

        self.dono = U.objects.create_user(
            email="donodocx@cert.test", password=SENHA, is_organizador=True
        )
        self.evento = _evento(self.dono)
        self.config = ConfiguracaoCertificado.objects.create(
            evento=self.evento, modo_layout="docx"
        )
        self.config.template_docx.save(
            "modelo.docx", ContentFile(_docx_template_bytes()), save=True
        )

    def test_render_docx_converte_e_usa_o_nome(self):
        from unittest import mock

        contexto = certificados.contexto_certificado(self.dono, evento=self.evento)
        capturado = {}

        def conversao_falsa(docx_path, outdir):
            from docx import Document

            capturado["texto"] = "\n".join(
                p.text for p in Document(docx_path).paragraphs
            )
            pdf = os.path.join(outdir, "certificado.pdf")
            with open(pdf, "wb") as arquivo:
                arquivo.write(b"%PDF-1.4 falso")
            return pdf

        with mock.patch.object(certificados, "_docx_para_pdf", side_effect=conversao_falsa):
            dados = certificados.render_docx(contexto, self.config)

        self.assertTrue(dados.startswith(b"%PDF"))
        self.assertIn(self.dono.get_full_name(), capturado["texto"])

    def test_render_pdf_despacha_para_docx(self):
        from unittest import mock

        contexto = certificados.contexto_certificado(self.dono, evento=self.evento)
        with mock.patch.object(certificados, "render_docx", return_value=b"%PDF-x") as chamada:
            dados = certificados.render_pdf(contexto, self.config)
        chamada.assert_called_once()
        self.assertEqual(dados, b"%PDF-x")

    def test_docx_sem_libreoffice_levanta_erro_claro(self):
        from unittest import mock

        contexto = certificados.contexto_certificado(self.dono, evento=self.evento)
        with mock.patch.object(
            certificados, "_docx_para_pdf", side_effect=FileNotFoundError("libreoffice")
        ):
            with self.assertRaises(certificados.CertificadoLayoutError):
                certificados.render_pdf(contexto, self.config)


class AtividadeOverrideTests(TestCase):
    @classmethod
    def setUpClass(cls):
        cls._media = tempfile.mkdtemp(prefix="media_ovr_")
        cls.enterClassContext(override_settings(MEDIA_ROOT=cls._media))
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(cls._media, ignore_errors=True)

    def setUp(self):
        self.dono = U.objects.create_user(
            email="donoovr@cert.test", password=SENHA, is_organizador=True
        )
        self.evento = _evento(self.dono)
        self.atividade = _atividade(self.evento)
        self.padrao = ConfiguracaoCertificado.objects.create(
            evento=self.evento, escopo="atividade", corpo="Padrão {{nome}}."
        )
        self.client.force_login(self.dono)

    def test_config_efetiva_usa_padrao_sem_override(self):
        self.assertEqual(certificados.config_efetiva(self.atividade), self.padrao)

    def test_criar_override_e_voltar_ao_padrao(self):
        url = reverse("organizador:certificado_atividade", args=[self.atividade.id])
        r = self.client.post(url, {"acao": "criar"})
        self.assertEqual(r.status_code, 302)
        override = ConfiguracaoCertificado.objects.get(atividade=self.atividade)
        self.assertEqual(override.escopo, "atividade")
        self.assertEqual(certificados.config_efetiva(self.atividade), override)

        r = self.client.post(
            reverse("organizador:certificado_atividade_usar_padrao", args=[self.atividade.id])
        )
        self.assertEqual(r.status_code, 302)
        self.assertFalse(
            ConfiguracaoCertificado.objects.filter(atividade=self.atividade).exists()
        )
        self.assertEqual(certificados.config_efetiva(self.atividade), self.padrao)

    def test_tela_padrao_das_atividades_lista(self):
        r = self.client.get(
            reverse("organizador:certificado_config", args=[self.evento.id, "atividades"])
        )
        self.assertEqual(r.status_code, 200)
        self.assertIn(self.atividade.titulo, r.content.decode("utf-8", "ignore"))


class CargaHorariaTests(TestCase):
    def setUp(self):
        self.dono = U.objects.create_user(
            email="carga@cert.test", password=SENHA, is_organizador=True
        )
        self.evento = _evento(self.dono)

    def test_carga_calculada_pelo_horario(self):
        a = _atividade(self.evento)  # 10:00–11:00
        ctx = certificados.contexto_certificado(self.dono, atividade=a, evento=self.evento)
        self.assertEqual(ctx["carga_horaria"], "1h")

    def test_carga_com_minutos(self):
        a = _atividade(
            self.evento,
            data_hora_inicio=datetime(2026, 10, 10, 10, 0, tzinfo=tz.utc),
            data_hora_fim=datetime(2026, 10, 10, 11, 30, tzinfo=tz.utc),
        )
        ctx = certificados.contexto_certificado(self.dono, atividade=a, evento=self.evento)
        self.assertEqual(ctx["carga_horaria"], "1h30")

    def test_campo_explicito_vence(self):
        a = _atividade(self.evento, carga_horaria=4)
        ctx = certificados.contexto_certificado(self.dono, atividade=a, evento=self.evento)
        self.assertEqual(ctx["carga_horaria"], "4h")

    def test_evento_sem_carga_horaria(self):
        ctx = certificados.contexto_certificado(self.dono, evento=self.evento)
        self.assertEqual(ctx["carga_horaria"], "")


class PercentualTests(TestCase):
    def setUp(self):
        self.dono = U.objects.create_user(
            email="perc@cert.test", password=SENHA, is_organizador=True
        )
        self.evento = _evento(self.dono)
        self.pessoa = U.objects.create_user(email="alunoP@cert.test", password=SENHA)

    def test_percentual_atingido_e_minimo(self):
        a1 = _atividade(self.evento, titulo="A1")
        a2 = _atividade(self.evento, titulo="A2")
        _atividade(self.evento, titulo="A3")
        _atividade(self.evento, titulo="A4")
        Presenca.objects.create(atividade=a1, participante=self.pessoa)
        Presenca.objects.create(atividade=a2, participante=self.pessoa)
        self.assertEqual(certificados.percentual_participacao(self.pessoa, self.evento), 50)
        ctx = certificados.contexto_certificado(self.pessoa, evento=self.evento)
        self.assertEqual(ctx["percentual_participacao"], "50%")
        self.assertEqual(ctx["percentual_minimo"], "75%")

    def test_atividade_sem_percentual(self):
        a = _atividade(self.evento)
        ctx = certificados.contexto_certificado(self.pessoa, atividade=a, evento=self.evento)
        self.assertEqual(ctx["percentual_participacao"], "")
        self.assertEqual(ctx["percentual_minimo"], "")


class FormEscopoTests(TestCase):
    def test_campos_por_escopo(self):
        from eventos.forms import ConfiguracaoCertificadoForm

        evento = ConfiguracaoCertificadoForm(escopo="evento")
        self.assertNotIn("carga_horaria_padrao", evento.fields)
        self.assertIn("percentual", evento.fields)

        atividade = ConfiguracaoCertificadoForm(escopo="atividade")
        self.assertIn("carga_horaria_padrao", atividade.fields)
        self.assertNotIn("percentual", atividade.fields)


class CertificadoStrTests(TestCase):
    def test_str_sem_atividade_e_evento_nao_quebra(self):
        dono = U.objects.create_user(email="str@cert.test", password=SENHA)
        certificado = Certificado.objects.create(participante=dono)
        self.assertIn("Certificado de", str(certificado))

    def test_str_com_atividade(self):
        dono = U.objects.create_user(
            email="str2@cert.test", password=SENHA, is_organizador=True
        )
        evento = _evento(dono)
        atividade = _atividade(evento)
        certificado = Certificado.objects.create(participante=dono, atividade=atividade)
        self.assertIn(atividade.titulo, str(certificado))

    def test_str_com_evento(self):
        dono = U.objects.create_user(
            email="str3@cert.test", password=SENHA, is_organizador=True
        )
        evento = _evento(dono)
        certificado = Certificado.objects.create(participante=dono, evento=evento)
        self.assertIn(evento.title, str(certificado))


class FundoTests(TestCase):
    """No modo texto, a imagem de fundo é opcional: usa se existir, senão padrão."""

    def test_com_imagem_desenha_o_fundo(self):
        from unittest import mock

        config = mock.Mock()
        config.layout_fundo = mock.Mock()
        config.layout_fundo.path = "/tmp/fundo.png"
        canvas = mock.Mock()
        with mock.patch.object(certificados, "ImageReader"):
            certificados._desenhar_fundo(canvas, config, 842, 595)
        canvas.drawImage.assert_called_once()

    def test_sem_imagem_usa_o_padrao(self):
        from unittest import mock

        config = mock.Mock()
        config.layout_fundo = None
        canvas = mock.Mock()
        certificados._desenhar_fundo(canvas, config, 842, 595)
        canvas.drawImage.assert_not_called()
        canvas.rect.assert_called()


class EncerramentoTests(TestCase):
    """O botão de emitir só aparece após o encerramento (regra da UI)."""

    def setUp(self):
        self.dono = U.objects.create_user(
            email="dono@enc.test", password=SENHA, is_organizador=True
        )

    def test_evento_encerrado(self):
        passado = _evento(
            self.dono, data_inicio=date(2020, 1, 1), data_fim=date(2020, 1, 2)
        )
        futuro = _evento(
            self.dono, data_inicio=date(2099, 1, 1), data_fim=date(2099, 1, 2)
        )
        self.assertTrue(passado.encerrado)
        self.assertFalse(futuro.encerrado)

    def test_atividade_encerrada(self):
        evento = _evento(self.dono)
        passada = _atividade(
            evento, data_hora_fim=datetime(2020, 1, 1, 10, 0, tzinfo=tz.utc)
        )
        futura = _atividade(
            evento, data_hora_fim=datetime(2099, 1, 1, 10, 0, tzinfo=tz.utc)
        )
        self.assertTrue(passada.encerrada)
        self.assertFalse(futura.encerrada)
