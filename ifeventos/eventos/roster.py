"""Pré-carga de pessoas (planilha) e completamento do perfil no 1º acesso.

Um ÚNICO arquivo (CSV, XLS ou XLSX) alimenta a tabela `PessoaRoster`, com a
chave sendo o e-mail PESSOAL. O mesmo arquivo serve para TODOS os vínculos:
além de `email`, `nome` e `cpf`, as colunas são as chaves de
`settings.METADADOS_PARTICIPANTE` (vinculo, matricula, curso, turma, ano,
funcao…). O que é específico do vínculo fica em `dados` (JSON), validado
pelo MESMO schema do formulário.

No primeiro acesso (Google ou cadastro local), `completar_do_roster` preenche
metadados/CPF/nome que a pessoa ainda não informou e marca a linha (usada,
confere, o que foi salvo). Roda UMA vez, na criação da conta.

Regras:
  - chave = e-mail; colunas fora do schema são ignoradas com aviso;
  - a validação é POR VÍNCULO (Aluno exige matrícula/curso; Servidor exige
    função; os demais só o vínculo);
  - e-mail fora da planilha, ou conta que já existia, => nada acontece.
"""

import logging
import os
import unicodedata

from django.utils import timezone

from . import metadados
from .models import PessoaRoster
from .validators import cpf_e_valido

logger = logging.getLogger(__name__)

# Colunas próprias da planilha (não são metadados).
COL_EMAIL = "email"
COL_NOME = "nome"
COL_CPF = "cpf"
FIXAS = (COL_EMAIL, COL_NOME, COL_CPF)


def _chave(texto):
    """Nome de coluna/chave normalizado: minúsculas, sem acento, sem pontas."""
    base = unicodedata.normalize("NFD", str(texto or "").lower())
    return "".join(c for c in base if unicodedata.category(c) != "Mn").strip()


def _texto(valor):
    """Valor de célula como texto. Números inteiros do Excel perdem o `.0`."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def cabecalhos():
    """Cabeçalho do arquivo: `email`, `nome`, `cpf` + as chaves de metadados."""
    return list(FIXAS) + [c["chave"] for c in metadados.campos()]


# ---------------------------------------------------------------------------
# Leitura do arquivo (CSV / XLS / XLSX)
# ---------------------------------------------------------------------------

def ler_csv(caminho):
    """Lê um CSV (utf-8-sig; delimitador , ; ou tab) e devolve as linhas."""
    from .importacao import ler_linhas

    with open(caminho, "rb") as arquivo:
        _cabecalho, linhas = ler_linhas(arquivo)
    return linhas


def ler_xls(caminho):
    """Lê a 1ª aba de um `.xls` (xlrd) e devolve as linhas."""
    import xlrd

    livro = xlrd.open_workbook(caminho)
    aba = livro.sheet_by_index(0)
    if aba.nrows == 0:
        return []
    cabecalho = [_chave(v) for v in aba.row_values(0)]
    linhas = []
    for r in range(1, aba.nrows):
        valores = aba.row_values(r)
        linhas.append({
            cabecalho[i]: _texto(valores[i])
            for i in range(len(cabecalho)) if i < len(valores)
        })
    return linhas


def ler_xlsx(caminho):
    """Lê a aba ativa de um `.xlsx` (openpyxl) e devolve as linhas."""
    from openpyxl import load_workbook

    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        linhas_iter = livro.active.iter_rows(values_only=True)
        try:
            cabecalho = [_chave(v) for v in next(linhas_iter)]
        except StopIteration:
            return []
        linhas = []
        for valores in linhas_iter:
            linhas.append({
                cabecalho[i]: _texto(valores[i])
                for i in range(len(cabecalho)) if i < len(valores)
            })
        return linhas
    finally:
        livro.close()


def ler_arquivo(caminho):
    """Escolhe o leitor pela extensão (.csv / .xls / .xlsx)."""
    ext = os.path.splitext(caminho)[1].lower()
    if ext == ".csv":
        return ler_csv(caminho)
    if ext == ".xls":
        return ler_xls(caminho)
    if ext in (".xlsx", ".xlsm"):
        return ler_xlsx(caminho)
    raise ValueError(
        f"Extensão não suportada: {ext or '(sem extensão)'}. Use .csv, .xls ou .xlsx."
    )


# ---------------------------------------------------------------------------
# Importação
# ---------------------------------------------------------------------------

def importar_linhas(linhas):
    """Núcleo: lista de linhas (dict) -> relatório por linha.

    A validação é por vínculo (`metadados.validar`), então o mesmo arquivo
    aceita Aluno, Servidor, Colaborador, Estagiário e Comunidade externa.
    """
    chaves = [c["chave"] for c in metadados.campos()]
    chaves_validas = set(FIXAS) | set(chaves)
    relatorio = {"total": len(linhas or []), "importados": 0, "atualizados": 0,
                 "ignorados": 0, "erros": 0, "linhas": []}

    for indice, bruta in enumerate(linhas or [], start=1):
        linha = {_chave(k): v for k, v in (bruta or {}).items()}
        email = _texto(linha.get(COL_EMAIL)).lower()
        item = {"indice": indice, "email": email, "status": None,
                "motivo": "", "avisos": []}

        def terminar(status, motivo=""):
            item["status"] = status
            item["motivo"] = motivo
            relatorio[status] = relatorio.get(status, 0) + 1
            relatorio["linhas"].append(item)

        if not email:
            terminar("erros", "linha sem e-mail.")
            continue

        for extra in sorted(set(linha) - chaves_validas):
            item["avisos"].append(f"coluna '{extra}' ignorada.")

        valores = {c: _texto(linha.get(c)) for c in chaves}
        for campo in metadados.campos():
            if metadados.visivel(campo, valores):
                continue
            if valores.get(campo["chave"]):
                item["avisos"].append(
                    f"'{campo['rotulo']}' ignorado (não se aplica ao vínculo)."
                )
        dados, erros = metadados.limpar(valores)
        if erros:
            mensagens = [m for lista in erros.values() for m in lista]
            terminar("erros", "; ".join(mensagens))
            continue

        cpf = _texto(linha.get(COL_CPF))
        if cpf and not cpf_e_valido(cpf):
            terminar("erros", "CPF inválido.")
            continue

        _obj, criado = PessoaRoster.objects.update_or_create(
            email=email,
            defaults={
                "nome": _texto(linha.get(COL_NOME)),
                "cpf": cpf,
                "dados": dados,
                "vinculo": dados.get("vinculo", ""),
            },
        )
        terminar("importados" if criado else "atualizados")

    return relatorio


def importar_arquivo(caminho):
    """Lê o arquivo (csv/xls/xlsx) e importa para `PessoaRoster`."""
    return importar_linhas(ler_arquivo(caminho))


# ---------------------------------------------------------------------------
# Completamento do perfil (1º acesso)
# ---------------------------------------------------------------------------

def dividir_nome(nome):
    """Divide um nome completo em ``(primeiro, sobrenome)``.

    "Maria da Silva" -> ("Maria", "da Silva"); "Madonna" -> ("Madonna", "").
    Cada parte é truncada em 150 (limite do model `Participante`).
    """
    partes = (nome or "").split()
    if not partes:
        return "", ""
    return partes[0][:150], " ".join(partes[1:])[:150]


def _conferem(linha, dados, cpf):
    """O que a pessoa salvou bate com a planilha? (metadados + CPF; nome fora).

    O nome fica de fora de propósito: no Google ele vem do Google, não da
    planilha, então comparar daria divergência que não é erro.
    """
    for campo in metadados.campos():
        chave = campo["chave"]
        atual = str((dados or {}).get(chave, "") or "").strip()
        planilha = str((linha.dados or {}).get(chave, "") or "").strip()
        if atual != planilha:
            return False
    return (cpf or "") == (linha.cpf or "")


def completar_do_roster(participante):
    """Preenche metadados/CPF/nome do participante a partir do `PessoaRoster`.

    Roda só na criação da conta (ver `signals.py`). Não faz nada se o e-mail
    não estiver na planilha. O que a própria pessoa já informou tem prioridade.
    Marca a linha como usada e se o que a pessoa salvou confere com a planilha.
    Nunca levanta exceção — o login não pode quebrar por causa disto.
    """
    try:
        if participante is None or not getattr(participante, "pk", None):
            return False
        email = (participante.email or "").strip().lower()
        if not email:
            return False

        linha = PessoaRoster.objects.filter(email=email).first()
        if linha is None:
            return False

        atual = metadados.dados_de(participante)
        base = dict(linha.dados or {})
        base.update(atual)  # o que a pessoa já informou manda

        dados, erros = metadados.limpar(base)
        if erros:
            logger.warning(
                "roster: metadados invalidos para %s: %s", email, erros
            )
            return False

        metadados.salvar(participante, dados)

        campos = []
        if linha.cpf and not participante.cpf:
            participante.cpf = linha.cpf
            campos.append("cpf")
        primeiro, sobrenome = dividir_nome(linha.nome)
        if primeiro and not participante.first_name:
            participante.first_name = primeiro
            campos.append("first_name")
        if sobrenome and not participante.last_name:
            participante.last_name = sobrenome
            campos.append("last_name")
        if campos:
            participante.save(update_fields=campos)

        # Rastreio do uso: quando foi usada, se confere e o que a pessoa salvou.
        linha.usado_em = timezone.now()
        linha.confere = _conferem(linha, dados, participante.cpf)
        linha.dados_usuario = {
            "nome": f"{participante.first_name} {participante.last_name}".strip(),
            "cpf": participante.cpf,
            "metadados": dados,
        }
        linha.save(update_fields=["usado_em", "confere", "dados_usuario"])
        return True
    except Exception:
        logger.exception("roster: falha ao completar perfil de %r", getattr(participante, "email", None))
        return False
