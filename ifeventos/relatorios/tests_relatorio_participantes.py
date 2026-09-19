"""Testes do relatório por participante (F1/F3), da permissão (F0) e da IA dos gráficos."""

from datetime import timedelta
import json
from unittest import mock
from unittest.mock import AsyncMock

from django.contrib.auth import get_user_model
from django.test import TestCase, TransactionTestCase
from django.urls import reverse
from django.utils import timezone
from asgiref.sync import async_to_sync

from eventos.models import (
    Atividade, Certificado, Evento, Inscricao, Presenca, TipoAtividade,
)
from relatorios import agente_graficos, agregacoes

U = get_user_model()
SENHA = "SenhaForte123!"


def _resposta_fake(conteudo):
    import json
    resposta = mock.MagicMock()
    resposta.choices = [mock.MagicMock()]
    resposta.choices[0].message.content = json.dumps(conteudo)
    resposta.usage = None
    return resposta


class BaseRelatorioTests(TestCase):
    def setUp(self):
        self.org = U.objects.create_user(
            email="org_rel@example.com", password=SENHA, cpf="12345678909",
            is_organizador=True,
        )
        self.outro = U.objects.create_user(
            email="outro_rel@example.com", password=SENHA, cpf="39053344705",
            is_organizador=True,
        )
        self.participante = U.objects.create_user(
            email="part_rel@example.com", password=SENHA, cpf="11144477735",
        )
        self.tipo = TipoAtividade.objects.create(nome="Oficina")
        hoje = timezone.localdate()
        self.evento = Evento.objects.create(
            title="Evento Relatório", description="d", local="Campus",
            data_inicio=hoje, data_fim=hoje + timedelta(days=1), organizador=self.org,
        )
        self.atividade = Atividade.objects.create(
            evento=self.evento, titulo="Oficina A", descricao="d", tipo=self.tipo,
            n_vagas=10, data_hora_inicio=timezone.now(),
            data_hora_fim=timezone.now() + timedelta(hours=2),
        )
        Inscricao.objects.create(participante=self.participante, atividade=self.atividade)
        Presenca.objects.create(
            atividade=self.atividade, participante=self.participante, papel="participante",
        )
        Certificado.objects.create(
            participante=self.participante, atividade=self.atividade, evento=self.evento,
        )

    def url(self, nome, *args):
        return reverse("organizador:%s" % nome, args=list(args))


class PermissaoRelatoriosTests(BaseRelatorioTests):
    """F0: os relatórios existentes deixam de ser abertos a qualquer login."""

    def test_participante_recebe_403(self):
        self.client.force_login(self.participante)
        resposta = self.client.get(self.url("relatorio_inscricoes", self.evento.id))
        self.assertEqual(resposta.status_code, 403)

    def test_dono_recebe_200(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("relatorio_inscricoes", self.evento.id))
        self.assertEqual(resposta.status_code, 200)

    def test_outro_organizador_recebe_200(self):
        self.client.force_login(self.outro)
        resposta = self.client.get(self.url("relatorio_inscricoes", self.evento.id))
        self.assertEqual(resposta.status_code, 200)

    def test_anonimo_redireciona(self):
        resposta = self.client.get(self.url("relatorio_inscricoes", self.evento.id))
        self.assertEqual(resposta.status_code, 302)

    def test_lista_presenca_protegida(self):
        self.client.force_login(self.participante)
        resposta = self.client.get(
            self.url("relatorio_lista_presenca", self.atividade.id)
        )
        self.assertEqual(resposta.status_code, 403)


class AgregacoesTests(BaseRelatorioTests):
    def test_resumo_por_pessoa(self):
        linhas = agregacoes.resumo_por_pessoa(self.evento)
        linha = next(l for l in linhas if l["pessoa"].id == self.participante.id)
        self.assertEqual(linha["n_inscricoes"], 1)
        self.assertEqual(linha["n_presencas"], 1)
        self.assertEqual(linha["n_certificados"], 1)
        self.assertEqual(linha["carga_horaria"], 2.0)

    def test_grupo_invariante(self):
        linhas = agregacoes.resumo_por_pessoa(self.evento)
        grupos = agregacoes.resumo_por_grupo(linhas)
        self.assertEqual(sum(g["pessoas"] for g in grupos), len(linhas))
        self.assertEqual(
            sum(g["presencas"] for g in grupos),
            sum(l["n_presencas"] for l in linhas),
        )

    def test_filtro_busca(self):
        self.assertEqual(len(agregacoes.resumo_por_pessoa(self.evento, termo="Part")), 1)
        self.assertEqual(len(agregacoes.resumo_por_pessoa(self.evento, termo="zzz")), 0)


class RelatorioParticipantesViewTests(BaseRelatorioTests):
    def test_pagina_renderiza(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("relatorio_participantes", self.evento.id))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Relatório por participante")
        self.assertContains(resposta, self.participante.email)

    def test_detalhe_renderiza(self):
        self.client.force_login(self.org)
        resposta = self.client.get(
            self.url("relatorio_participante_detalhe", self.evento.id, self.participante.id)
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, self.participante.email)

    def test_export_csv(self):
        self.client.force_login(self.org)
        resposta = self.client.get(
            self.url("relatorio_participantes", self.evento.id) + "?export=csv"
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("text/csv", resposta["Content-Type"])


class AgenteGraficosTests(BaseRelatorioTests, TransactionTestCase):
    def test_curadoria_fallback(self):
        linhas = agregacoes.resumo_por_pessoa(self.evento)
        disponiveis = agregacoes.graficos_participantes(linhas)
        with mock.patch("eventos.services.get_openai_client", side_effect=RuntimeError("x")):
            resultado = async_to_sync(agente_graficos.curar)(
                self.evento, disponiveis, {"pessoas": len(linhas)}
            )
        self.assertEqual(resultado["origem"], "heuristica")
        self.assertTrue(resultado["ids"])

    def test_curadoria_ia_descarta_id_invalido(self):
        disponiveis = [{"id": "participantes_vinculo", "titulo": "Vínculo"}]
        fake = mock.MagicMock()
        fake.chat = mock.MagicMock()
        fake.chat.completions.create = AsyncMock(
            return_value=_resposta_fake({"ids": ["inexistente", "participantes_vinculo"], "legenda": "ok"})
        )
        with mock.patch("eventos.services.get_openai_client", return_value=fake):
            resultado = async_to_sync(agente_graficos.curar)(
                self.evento, disponiveis, {"pessoas": 1}
            )
        self.assertEqual(resultado["ids"], ["participantes_vinculo"])

    def test_endpoint_curadoria_exige_organizador(self):
        self.client.force_login(self.participante)
        resposta = self.client.post(
            self.url("graficos_curadoria", self.evento.id)
        )
        self.assertEqual(resposta.status_code, 403)

    def test_grafico_por_descricao_fallback_casa_palavra(self):
        disponiveis = [
            {"id": "participantes_carga", "titulo": "Carga horária por participante"},
            {"id": "participantes_presenca", "titulo": "Presenças × ausências"},
        ]
        with mock.patch("eventos.services.get_openai_client", side_effect=RuntimeError("x")):
            resultado = async_to_sync(agente_graficos.grafico_por_descricao)(
                self.evento, disponiveis, "carga horária dos participantes"
            )
        self.assertEqual(resultado["origem"], "heuristica")
        self.assertEqual(resultado["id"], "participantes_carga")

    def test_grafico_por_descricao_fallback_sem_casa_usa_primeiro(self):
        disponiveis = [
            {"id": "a", "titulo": "Gráfico A"},
            {"id": "b", "titulo": "Gráfico B"},
        ]
        with mock.patch("eventos.services.get_openai_client", side_effect=RuntimeError("x")):
            resultado = async_to_sync(agente_graficos.grafico_por_descricao)(
                self.evento, disponiveis, "xyz"
            )
        self.assertEqual(resultado["id"], "a")

    def test_grafico_por_descricao_ia_usa_id_valido(self):
        disponiveis = [
            {"id": "participantes_presenca", "titulo": "Presenças × ausências"},
            {"id": "participantes_carga", "titulo": "Carga horária por participante"},
        ]
        fake = mock.MagicMock()
        fake.chat = mock.MagicMock()
        fake.chat.completions.create = AsyncMock(
            return_value=_resposta_fake({"id": "participantes_presenca", "legenda": "presença"})
        )
        with mock.patch("eventos.services.get_openai_client", return_value=fake):
            resultado = async_to_sync(agente_graficos.grafico_por_descricao)(
                self.evento, disponiveis, "quem foi e quem não foi"
            )
        self.assertEqual(resultado["origem"], "ia")
        self.assertEqual(resultado["id"], "participantes_presenca")

    def test_endpoint_grafico_por_descricao_organizador(self):
        self.client.force_login(self.org)
        fake = mock.MagicMock()
        fake.chat = mock.MagicMock()
        fake.chat.completions.create = AsyncMock(
            return_value=_resposta_fake({"id": "participantes_carga", "legenda": "carga"})
        )
        with mock.patch("eventos.services.get_openai_client", return_value=fake):
            resposta = self.client.post(
                self.url("grafico_por_descricao", self.evento.id),
                data=json.dumps({"texto": "carga horária"}),
                content_type="application/json",
            )
        self.assertEqual(resposta.status_code, 200)
        dados = resposta.json()
        self.assertEqual(dados["id"], "participantes_carga")
        self.assertEqual(dados["origem"], "ia")
        self.assertEqual(dados["grafico"]["id"], "participantes_carga")
        self.assertTrue(dados["grafico"]["series"])

    def test_grafico_por_descricao_pode_retornar_grafico_do_evento(self):
        # O catálogo unificado inclui os gráficos do evento (ex.: ocupação por
        # sala) que a página pode não mostrar — o endpoint devolve o gráfico
        # montado para o frontend anexar como card novo.
        self.client.force_login(self.org)
        fake = mock.MagicMock()
        fake.chat = mock.MagicMock()
        fake.chat.completions.create = AsyncMock(
            return_value=_resposta_fake({"id": "ocupacao_sala", "legenda": "por sala"})
        )
        with mock.patch("eventos.services.get_openai_client", return_value=fake):
            resposta = self.client.post(
                self.url("grafico_por_descricao", self.evento.id),
                data=json.dumps({"pagina": "oficinas", "texto": "ocupação por sala"}),
                content_type="application/json",
            )
        self.assertEqual(resposta.status_code, 200)
        dados = resposta.json()
        self.assertEqual(dados["id"], "ocupacao_sala")
        self.assertEqual(dados["grafico"]["id"], "ocupacao_sala")
        self.assertIn("labels", dados["grafico"])

    def test_endpoint_grafico_por_descricao_exige_organizador(self):
        self.client.force_login(self.participante)
        resposta = self.client.post(
            self.url("grafico_por_descricao", self.evento.id),
            data=json.dumps({"texto": "qualquer"}),
            content_type="application/json",
        )
        self.assertEqual(resposta.status_code, 403)


class RelatorioTurmasTests(BaseRelatorioTests):
    def setUp(self):
        super().setUp()
        from eventos.models import ParticipanteMetadados

        ParticipanteMetadados.objects.update_or_create(
            participante=self.participante,
            defaults={"dados": {"vinculo": "Aluno", "curso": "Informática",
                                "turma": "Turma 1", "matricula": "12345"}},
        )

    def test_pagina_renderiza_e_grupo(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("relatorio_turmas", self.evento.id))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Relatório por turma")
        self.assertContains(resposta, "Informática · Turma 1")

    def test_invariante_soma_dos_grupos(self):
        linhas = agregacoes.resumo_por_pessoa(self.evento)
        grupos = agregacoes.resumo_por_grupo(linhas)
        self.assertEqual(
            sum(g["pessoas"] for g in grupos), len(linhas)
        )
        self.assertEqual(
            sum(g["presencas"] for g in grupos),
            sum(l["n_presencas"] for l in linhas),
        )

    def test_grupo_de(self):
        linha = next(
            l for l in agregacoes.resumo_por_pessoa(self.evento)
            if l["pessoa"].id == self.participante.id
        )
        self.assertEqual(agregacoes.grupo_de(linha, "curso_turma_ano"), "Informática · Turma 1")
        self.assertEqual(agregacoes.grupo_de(linha, "curso"), "Informática")

    def test_drill_down_filtra_participantes(self):
        self.client.force_login(self.org)
        resposta = self.client.get(
            self.url("relatorio_participantes", self.evento.id)
            + "?agrupar=curso_turma_ano&grupo="
            + __import__("urllib.parse", fromlist=["quote"]).quote("Informática · Turma 1")
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, self.participante.email)

    def test_export_xlsx_tem_duas_abas(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("relatorio_turmas", self.evento.id) + "?export=xlsx")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("spreadsheetml", resposta["Content-Type"])

        from openpyxl import load_workbook
        from io import BytesIO

        workbook = load_workbook(BytesIO(resposta.content))
        self.assertEqual(workbook.sheetnames, ["Resumo", "Detalhe"])


class RelatorioOficinasTests(BaseRelatorioTests):
    """F4: relatório por tipo de atividade."""

    def test_pagina_renderiza(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("relatorio_oficinas", self.evento.id))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Oficina A")
        self.assertContains(resposta, "Oficina")

    def test_filtro_por_tipo(self):
        outro_tipo = TipoAtividade.objects.create(nome="Palestra")
        Atividade.objects.create(
            evento=self.evento, titulo="Palestra X", descricao="d", tipo=outro_tipo,
            n_vagas=5, data_hora_inicio=timezone.now(),
            data_hora_fim=timezone.now() + timedelta(hours=1),
        )
        self.client.force_login(self.org)
        resposta = self.client.get(
            self.url("relatorio_oficinas", self.evento.id) + f"?tipo={self.tipo.id}"
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "Oficina A")
        self.assertNotContains(resposta, "Palestra X")

    def test_export_xlsx(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("relatorio_oficinas", self.evento.id) + "?export=xlsx")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("spreadsheetml", resposta["Content-Type"])

        from openpyxl import load_workbook
        from io import BytesIO

        workbook = load_workbook(BytesIO(resposta.content))
        self.assertEqual(workbook.sheetnames, ["Resumo"])

    def test_participante_recebe_403(self):
        self.client.force_login(self.participante)
        resposta = self.client.get(self.url("relatorio_oficinas", self.evento.id))
        self.assertEqual(resposta.status_code, 403)


class BadgeOutroOrganizadorTests(BaseRelatorioTests):
    """R3: badge 'evento de outro organizador' nas telas do organizador."""

    def test_badge_aparece_para_outro_organizador(self):
        self.client.force_login(self.outro)
        resposta = self.client.get(self.url("atividades_evento", self.evento.id))
        self.assertEqual(resposta.status_code, 200)
        self.assertContains(resposta, "outro organizador")

    def test_badge_ausente_para_dono(self):
        self.client.force_login(self.org)
        resposta = self.client.get(self.url("atividades_evento", self.evento.id))
        self.assertEqual(resposta.status_code, 200)
        self.assertNotContains(resposta, "outro organizador")
